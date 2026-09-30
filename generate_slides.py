import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation (16:9 Widescreen)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors
BG_COLOR = RGBColor(248, 250, 252)       # Slate 50
CARD_BG = RGBColor(255, 255, 255)        # White
CARD_BORDER = RGBColor(226, 232, 240)    # Slate 200
TEXT_MAIN = RGBColor(15, 23, 42)         # Slate 900
TEXT_MUTED = RGBColor(100, 116, 139)     # Slate 500
ACCENT_GREEN = RGBColor(16, 185, 129)    # Emerald 500
ACCENT_GREEN_BG = RGBColor(236, 253, 245)# Emerald 50
ACCENT_BLUE = RGBColor(37, 99, 235)      # Blue 600
ACCENT_BLUE_BG = RGBColor(239, 246, 255) # Blue 50
ACCENT_PURPLE = RGBColor(124, 58, 237)   # Purple 600
ACCENT_RED = RGBColor(220, 38, 38)       # Red 600
ACCENT_AMBER = RGBColor(217, 119, 6)     # Amber 600

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_header(slide, title, category="EXCEL FORMULAS & FUNCTIONS"):
    # Category Pill / Tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(10), Inches(0.4))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category.upper()
    p_tag.font.name = "Arial"
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_GREEN

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.7))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Arial"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN

def add_card(slide, left, top, width, height, bg_rgb=CARD_BG, border_rgb=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_rgb
    shape.line.color.rgb = border_rgb
    shape.line.width = Pt(1.5)
    return shape

# ==========================================
# SLIDE 1: Title Slide
# ==========================================
s1 = prs.slides.add_slide(blank_layout)
set_slide_background(s1)

# Large Title Card
c1 = add_card(s1, 1.2, 1.5, 10.933, 4.5, bg_rgb=CARD_BG, border_rgb=CARD_BORDER)

# Title content
tbox = s1.shapes.add_textbox(Inches(1.8), Inches(2.2), Inches(9.7), Inches(3.0))
tf = tbox.text_frame
tf.word_wrap = True

p1 = tf.paragraphs[0]
p1.text = "Excel Core Functions & Diagnostics"
p1.font.name = "Arial"
p1.font.size = Pt(38)
p1.font.bold = True
p1.font.color.rgb = TEXT_MAIN

p2 = tf.add_paragraph()
p2.text = "Pages 18 – 20 Comprehensive Course Guide"
p2.font.name = "Arial"
p2.font.size = Pt(20)
p2.font.color.rgb = ACCENT_GREEN
p2.font.bold = True
p2.space_before = Pt(14)

p3 = tf.add_paragraph()
p3.text = "Mathematical Formulas • Statistical Analysis • Rounding Rules • Error Troubleshooting"
p3.font.name = "Arial"
p3.font.size = Pt(14)
p3.font.color.rgb = TEXT_MUTED
p3.space_before = Pt(18)


# ==========================================
# SLIDE 2: Anatomy of a Function
# ==========================================
s2 = prs.slides.add_slide(blank_layout)
set_slide_background(s2)
add_header(s2, "Anatomy of an Excel Function", "Fundamentals")

# Left Column: Syntax & Rules
add_card(s2, 0.8, 1.7, 5.6, 5.2)
l_box = s2.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
ltf = l_box.text_frame
ltf.word_wrap = True

lp1 = ltf.paragraphs[0]
lp1.text = "What is a Function?"
lp1.font.size = Pt(18)
lp1.font.bold = True
lp1.font.color.rgb = TEXT_MAIN

lp2 = ltf.add_paragraph()
lp2.text = "A predefined formula built into Excel that performs specific calculations automatically."
lp2.font.size = Pt(13)
lp2.font.color.rgb = TEXT_MUTED
lp2.space_before = Pt(8)

lp3 = ltf.add_paragraph()
lp3.text = "3 Key Components:"
lp3.font.size = Pt(14)
lp3.font.bold = True
lp3.font.color.rgb = TEXT_MAIN
lp3.space_before = Pt(14)

lp4 = ltf.add_paragraph()
lp4.text = "1. Equal Sign (=)\n   Tells Excel to calculate rather than treat as text."
lp4.font.size = Pt(12)
lp4.font.color.rgb = TEXT_MUTED
lp4.space_before = Pt(6)

lp5 = ltf.add_paragraph()
lp5.text = "2. Function Name (e.g., SUM, AVERAGE)\n   Identifies the operation to execute."
lp5.font.size = Pt(12)
lp5.font.color.rgb = TEXT_MUTED
lp5.space_before = Pt(6)

lp6 = ltf.add_paragraph()
lp6.text = "3. Arguments inside Parentheses ( )\n   The cell values or ranges (A1:A10) being calculated."
lp6.font.size = Pt(12)
lp6.font.color.rgb = TEXT_MUTED
lp6.space_before = Pt(6)

# Right Column: Visual Syntax Breakdown
add_card(s2, 6.8, 1.7, 5.7, 5.2)
r_box = s2.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
rtf = r_box.text_frame
rtf.word_wrap = True

rp1 = rtf.paragraphs[0]
rp1.text = "Formula Syntax Breakdown"
rp1.font.size = Pt(18)
rp1.font.bold = True
rp1.font.color.rgb = TEXT_MAIN

# Code Pill inside right card
rp2 = rtf.add_paragraph()
rp2.text = "  = SUM ( A1 : A5 )  "
rp2.font.size = Pt(22)
rp2.font.bold = True
rp2.font.name = "Courier New"
rp2.font.color.rgb = ACCENT_BLUE
rp2.space_before = Pt(16)

rp3 = rtf.add_paragraph()
rp3.text = "Comparison: Manual vs. Function"
rp3.font.size = Pt(14)
rp3.font.bold = True
rp3.font.color.rgb = TEXT_MAIN
rp3.space_before = Pt(20)

rp4 = rtf.add_paragraph()
rp4.text = "❌ Manual:   = A1 + A2 + A3 + A4 + A5\n   (Tedious and breaks if rows are added)\n\n✅ Function:  = SUM(A1:A5)\n   (Fast, scalable, and dynamically updates)"
rp4.font.size = Pt(13)
rp4.font.color.rgb = TEXT_MUTED
rp4.space_before = Pt(8)


# ==========================================
# SLIDE 3: SUM & AVERAGE
# ==========================================
s3 = prs.slides.add_slide(blank_layout)
set_slide_background(s3)
add_header(s3, "1. SUM & 2. AVERAGE", "Aggregation Functions")

# Card 1: SUM
add_card(s3, 0.8, 1.7, 5.6, 5.2)
b1 = s3.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
t1 = b1.text_frame
t1.word_wrap = True

p = t1.paragraphs[0]
p.text = "SUM(range)"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

p = t1.add_paragraph()
p.text = "Adds all numerical values across selected cells."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(6)

p = t1.add_paragraph()
p.text = "Syntax: =SUM(number1, [number2], ...)"
p.font.size = Pt(13)
p.font.name = "Courier New"
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(12)

p = t1.add_paragraph()
p.text = "Example Dataset: [10, 20, 30, 40]\n=SUM(A1:A4)\n➜ Result: 100"
p.font.size = Pt(13)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(12)

p = t1.add_paragraph()
p.text = "💡 Tip: Ignores text and empty cells automatically."
p.font.size = Pt(12)
p.font.color.rgb = ACCENT_AMBER
p.font.italic = True
p.space_before = Pt(14)

# Card 2: AVERAGE
add_card(s3, 6.8, 1.7, 5.7, 5.2)
b2 = s3.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
t2 = b2.text_frame
t2.word_wrap = True

p = t2.paragraphs[0]
p.text = "AVERAGE(range)"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p = t2.add_paragraph()
p.text = "Calculates the arithmetic mean of numbers in a range."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(6)

p = t2.add_paragraph()
p.text = "Syntax: =AVERAGE(number1, [number2], ...)"
p.font.size = Pt(13)
p.font.name = "Courier New"
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(12)

p = t2.add_paragraph()
p.text = "Formula Logic:\nAverage = Total Sum ÷ Total Count\n\nExample: (10 + 20 + 30 + 40) ÷ 4\n➜ Result: 25"
p.font.size = Pt(13)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(12)


# ==========================================
# SLIDE 4: AVERAGE Deep Dive (Slide 19 Visualizer)
# ==========================================
s4 = prs.slides.add_slide(blank_layout)
set_slide_background(s4)
add_header(s4, "Visualizing AVERAGE: The Equal Redistribution Model", "Course Slide 19")

# Left Column: Concept Explanation
add_card(s4, 0.8, 1.7, 5.6, 5.2)
b_avg_l = s4.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
t_avg_l = b_avg_l.text_frame
t_avg_l.word_wrap = True

p = t_avg_l.paragraphs[0]
p.text = "The 'Leveling Water' Concept"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = TEXT_MAIN

p = t_avg_l.add_paragraph()
p.text = "Imagine 3 glasses with different water levels:\n• Glass 1 = 2 units\n• Glass 2 = 4 units\n• Glass 3 = 9 units"
p.font.size = Pt(13)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(10)

p = t_avg_l.add_paragraph()
p.text = "When poured together and shared equally:\nTotal = 2 + 4 + 9 = 15 units\nEach glass receives: 15 ÷ 3 = 5 units"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(12)

p = t_avg_l.add_paragraph()
p.text = "Key Insight:"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(14)

p = t_avg_l.add_paragraph()
p.text = "The AVERAGE is the single value where the sum of deviations above equals the sum of deviations below."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(4)

# Right Column: Visual Graphic Card
add_card(s4, 6.8, 1.7, 5.7, 5.2)
b_avg_r = s4.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
t_avg_r = b_avg_r.text_frame
t_avg_r.word_wrap = True

p = t_avg_r.paragraphs[0]
p.text = "Redistribution Simulation"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = TEXT_MAIN

p = t_avg_r.add_paragraph()
p.text = "Original Values:      [ 2 ]     [ 4 ]     [ 9 ]\n                       ⬇         ⬇         ⬇\nCombined Sum:                  15 Total\n                       ⬇         ⬇         ⬇\nBalanced Average:     [ 5 ]     [ 5 ]     [ 5 ]"
p.font.size = Pt(14)
p.font.name = "Courier New"
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p.space_before = Pt(16)

p = t_avg_r.add_paragraph()
p.text = "Comparison with MEDIAN:"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(20)

p = t_avg_r.add_paragraph()
p.text = "• AVERAGE = 5  (Calculated from all values)\n• MEDIAN  = 4  (The exact middle item)\n\nIf Glass 3 increases to 99:\n• AVERAGE jumps to 35 (Distorted by outlier!)\n• MEDIAN stays at 4 (Stable!)"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(6)


# ==========================================
# SLIDE 5: Rounding Family
# ==========================================
s5 = prs.slides.add_slide(blank_layout)
set_slide_background(s5)
add_header(s5, "3. ROUND, 4. ROUNDUP, 5. ROUNDDOWN & 6. INT", "Rounding Functions")

# 4 Quadrant Cards
cards_data = [
    ("ROUND(num, digits)", ACCENT_GREEN, "Standard Rounding Rule", "Rounds 0-4 DOWN, 5-9 UP.\n\n=ROUND(3.1415, 2)  ➜ 3.14\n=ROUND(3.1465, 2)  ➜ 3.15", 0.8, 1.7),
    ("ROUNDUP(num, digits)", ACCENT_BLUE, "Always Round Away from 0", "Forces rounding UP to next increment.\n\n=ROUNDUP(3.1415, 2) ➜ 3.15\n=ROUNDUP(12.01, 0)   ➜ 13", 6.8, 1.7),
    ("ROUNDDOWN(num, digits)", ACCENT_AMBER, "Always Round Toward 0", "Truncates / chops off extra digits.\n\n=ROUNDDOWN(3.1499, 2) ➜ 3.14\n=ROUNDDOWN(12.99, 0)   ➜ 12", 0.8, 4.4),
    ("INT(number)", ACCENT_PURPLE, "Nearest Lower Integer", "Drops all decimals, returns whole integer.\n\n=INT(7.9)   ➜ 7\n=INT(-3.2)  ➜ -4 (Rounds down!)", 6.8, 4.4),
]

for title_fn, color, subtitle, desc, left, top in cards_data:
    add_card(s5, left, top, 5.7, 2.5)
    box = s5.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.15), Inches(5.2), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title_fn
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = color
    
    p = tf.add_paragraph()
    p.text = subtitle
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(3)
    
    p = tf.add_paragraph()
    p.text = desc
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(4)


# ==========================================
# SLIDE 6: Extremes & Ranked Values
# ==========================================
s6 = prs.slides.add_slide(blank_layout)
set_slide_background(s6)
add_header(s6, "7. MAX, 8. MIN, 9. LARGE & 10. SMALL", "Extremes & Positional Functions")

# Left Column: MAX & MIN
add_card(s6, 0.8, 1.7, 5.6, 5.2)
b_ext_l = s6.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
t_ext_l = b_ext_l.text_frame
t_ext_l.word_wrap = True

p = t_ext_l.paragraphs[0]
p.text = "Absolute Extremes"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = TEXT_MAIN

p = t_ext_l.add_paragraph()
p.text = "=MAX(range)"
p.font.size = Pt(15)
p.font.bold = True
p.font.name = "Courier New"
p.font.color.rgb = ACCENT_GREEN
p.space_before = Pt(10)

p = t_ext_l.add_paragraph()
p.text = "Returns the highest value in a dataset.\nExample: =MAX(85, 92, 78, 99) ➜ 99"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(4)

p = t_ext_l.add_paragraph()
p.text = "=MIN(range)"
p.font.size = Pt(15)
p.font.bold = True
p.font.name = "Courier New"
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(16)

p = t_ext_l.add_paragraph()
p.text = "Returns the lowest value in a dataset.\nExample: =MIN(85, 92, 78, 99) ➜ 78"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(4)

# Right Column: LARGE & SMALL (k-th positions)
add_card(s6, 6.8, 1.7, 5.7, 5.2)
b_ext_r = s6.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
t_ext_r = b_ext_r.text_frame
t_ext_r.word_wrap = True

p = t_ext_r.paragraphs[0]
p.text = "Ranked k-th Extremes"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = TEXT_MAIN

p = t_ext_r.add_paragraph()
p.text = "=LARGE(range, k)"
p.font.size = Pt(15)
p.font.bold = True
p.font.name = "Courier New"
p.font.color.rgb = ACCENT_PURPLE
p.space_before = Pt(10)

p = t_ext_r.add_paragraph()
p.text = "Returns the k-th largest value.\n• k = 1 ➜ Same as MAX\n• k = 2 ➜ 2nd highest score\nExample: =LARGE({50, 80, 95, 70}, 2) ➜ 80"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(4)

p = t_ext_r.add_paragraph()
p.text = "=SMALL(range, k)"
p.font.size = Pt(15)
p.font.bold = True
p.font.name = "Courier New"
p.font.color.rgb = ACCENT_AMBER
p.space_before = Pt(16)

p = t_ext_r.add_paragraph()
p.text = "Returns the k-th smallest value.\n• k = 1 ➜ Same as MIN\n• k = 2 ➜ 2nd lowest score\nExample: =SMALL({50, 80, 95, 70}, 2) ➜ 70"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(4)


# ==========================================
# SLIDE 7: Statistics & Leaderboard
# ==========================================
s7 = prs.slides.add_slide(blank_layout)
set_slide_background(s7)
add_header(s7, "11. COUNT, 12. MEDIAN, 13. MODE & 14. RANK.EQ", "Statistical Distribution")

# 4 Cards Layout
stat_data = [
    ("COUNT(range)", ACCENT_GREEN, "Counts Numbers Only", "Only tallies cells with numeric values (ignores text and blanks).\n\nDataset: [25, 'Pass', 30, ' ']\n=COUNT(A1:A4) ➜ 2", 0.8, 1.7),
    ("MEDIAN(range)", ACCENT_BLUE, "The Exact Middle Value", "Finds the center of sorted data.\n\nDataset: [10, 20, 30, 40, 50]\n=MEDIAN(A1:A5) ➜ 30\n(Resistant to extreme outliers)", 6.8, 1.7),
    ("MODE.SNGL(range)", ACCENT_PURPLE, "Most Frequent Value", "Identifies the most repeated number.\n\nDataset: [5, 8, 8, 12, 15]\n=MODE.SNGL(A1:A5) ➜ 8", 0.8, 4.4),
    ("RANK.EQ(val, range, [order])", ACCENT_AMBER, "Leaderboard Rank Position", "Calculates a number's standing in a list.\n• order 0 (default) = High is 1st\n• order 1 = Lowest is 1st\n\n=RANK.EQ(95, A1:A10, 0) ➜ 1", 6.8, 4.4),
]

for title_fn, color, subtitle, desc, left, top in stat_data:
    add_card(s7, left, top, 5.7, 2.5)
    box = s7.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.15), Inches(5.2), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title_fn
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = color
    
    p = tf.add_paragraph()
    p.text = subtitle
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(3)
    
    p = tf.add_paragraph()
    p.text = desc
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(4)


# ==========================================
# SLIDE 8: Remainder Division (MOD)
# ==========================================
s8 = prs.slides.add_slide(blank_layout)
set_slide_background(s8)
add_header(s8, "15. MOD Function", "Mathematical Remainder")

# Left Column: Syntax & Logic
add_card(s8, 0.8, 1.7, 5.6, 5.2)
b_mod_l = s8.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
t_mod_l = b_mod_l.text_frame
t_mod_l.word_wrap = True

p = t_mod_l.paragraphs[0]
p.text = "MOD(number, divisor)"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = ACCENT_PURPLE

p = t_mod_l.add_paragraph()
p.text = "Returns the remainder after dividing a number by a divisor."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(6)

p = t_mod_l.add_paragraph()
p.text = "Mathematical Equation:"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(12)

p = t_mod_l.add_paragraph()
p.text = "MOD(n, d) = n - d * INT(n / d)"
p.font.size = Pt(13)
p.font.name = "Courier New"
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(4)

p = t_mod_l.add_paragraph()
p.text = "Examples:\n• =MOD(10, 3) ➜ 1  (10 ÷ 3 = 3 remainder 1)\n• =MOD(12, 4) ➜ 0  (Exact multiple!)\n• =MOD(17, 5) ➜ 2  (17 ÷ 5 = 3 remainder 2)"
p.font.size = Pt(13)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(14)

# Right Column: Real-World Use Cases
add_card(s8, 6.8, 1.7, 5.7, 5.2)
b_mod_r = s8.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
t_mod_r = b_mod_r.text_frame
t_mod_r.word_wrap = True

p = t_mod_r.paragraphs[0]
p.text = "Practical Excel Applications"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = TEXT_MAIN

p = t_mod_r.add_paragraph()
p.text = "1. Alternating Row Colors (Zebra Striping)\n   Formula: =MOD(ROW(), 2) = 0\n   Highlights every even row automatically."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(10)

p = t_mod_r.add_paragraph()
p.text = "2. Packaging / Leftover Inventory\n   If you have 125 eggs in cartons of 12:\n   • Full Cartons = INT(125 / 12) = 10\n   • Leftover Eggs = MOD(125, 12) = 5"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(10)

p = t_mod_r.add_paragraph()
p.text = "3. Time & Schedule Shifts\n   Converting elapsed minutes into hours and remaining minutes."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(10)


# ==========================================
# SLIDE 9: Hash Errors Diagnostic Lab (Slide 20)
# ==========================================
s9 = prs.slides.add_slide(blank_layout)
set_slide_background(s9)
add_header(s9, "Excel Hash Errors Diagnostic Lab", "Course Slide 20 Troubleshooting")

# 4 Error Cards
err_data = [
    ("####", ACCENT_AMBER, "Column Width Too Narrow", "CAUSE:\nThe cell contents are wider than the column width, or date is negative.\n\nFIX:\nDouble-click the column boundary header to auto-fit width.", 0.8, 1.7),
    ("#DIV/0!", ACCENT_RED, "Division By Zero", "CAUSE:\nA formula attempts to divide a number by 0 or by an empty cell.\n\nFIX:\nCheck denominator cell or wrap in: =IFERROR(A1/B1, 0)", 6.8, 1.7),
    ("#VALUE!", ACCENT_RED, "Wrong Data Type", "CAUSE:\nMathematical operator applied to text (e.g. =A1 + 'Apples').\n\nFIX:\nEnsure all operands contain valid numbers, or use =SUM() which ignores text.", 0.8, 4.4),
    ("#NAME?", ACCENT_PURPLE, "Misspelled Formula / Name", "CAUSE:\nExcel does not recognize the formula name (e.g. =SUMM(A1:A5)) or unquoted text.\n\nFIX:\nCheck spelling of function names and add quotes to text strings.", 6.8, 4.4),
]

for title_fn, color, subtitle, desc, left, top in err_data:
    add_card(s9, left, top, 5.7, 2.5)
    box = s9.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.15), Inches(5.2), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title_fn
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.name = "Courier New"
    p.font.color.rgb = color
    
    p = tf.add_paragraph()
    p.text = subtitle
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    p.space_before = Pt(3)
    
    p = tf.add_paragraph()
    p.text = desc
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(4)


# ==========================================
# SLIDE 10: 15 Functions Master Cheat Sheet
# ==========================================
s10 = prs.slides.add_slide(blank_layout)
set_slide_background(s10)
add_header(s10, "15 Core Functions Master Cheat Sheet", "Summary & Quick Reference")

# Master Summary Table
rows = 16
cols = 3
left = Inches(0.8)
top = Inches(1.6)
width = Inches(11.733)
height = Inches(5.3)

table_shape = s10.shapes.add_table(rows, cols, left, top, width, height)
table = table_shape.table
table.columns[0].width = Inches(2.8)
table.columns[1].width = Inches(3.6)
table.columns[2].width = Inches(5.333)

headers = ["Function", "Syntax", "Purpose / Key Behavior"]
for i, h in enumerate(headers):
    cell = table.cell(0, i)
    cell.fill.solid()
    cell.fill.fore_color.rgb = ACCENT_GREEN_BG
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN

data_rows = [
    ("1. SUM", "=SUM(range)", "Adds all numbers in range"),
    ("2. AVERAGE", "=AVERAGE(range)", "Calculates arithmetic mean (Sum ÷ Count)"),
    ("3. ROUND", "=ROUND(number, num_digits)", "Standard rounding (≥5 up, <5 down)"),
    ("4. ROUNDUP", "=ROUNDUP(number, num_digits)", "Always rounds away from 0"),
    ("5. ROUNDDOWN", "=ROUNDDOWN(number, num_digits)", "Always rounds toward 0 (truncates)"),
    ("6. INT", "=INT(number)", "Rounds down to nearest integer"),
    ("7. MAX", "=MAX(range)", "Returns highest value in range"),
    ("8. MIN", "=MIN(range)", "Returns lowest value in range"),
    ("9. LARGE", "=LARGE(range, k)", "Returns k-th largest value in list"),
    ("10. SMALL", "=SMALL(range, k)", "Returns k-th smallest value in list"),
    ("11. COUNT", "=COUNT(range)", "Counts cells containing numbers only"),
    ("12. MEDIAN", "=MEDIAN(range)", "Finds middle sorted value (outlier-proof)"),
    ("13. MODE.SNGL", "=MODE.SNGL(range)", "Finds most frequently occurring number"),
    ("14. RANK.EQ", "=RANK.EQ(val, ref, [order])", "Calculates leaderboard ranking position"),
    ("15. MOD", "=MOD(number, divisor)", "Returns remainder after division"),
]

for row_idx, row_data in enumerate(data_rows, start=1):
    for col_idx, text in enumerate(row_data):
        cell = table.cell(row_idx, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG
        p = cell.text_frame.paragraphs[0]
        p.text = text
        p.font.name = "Arial" if col_idx != 1 else "Courier New"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MAIN if col_idx == 0 else (ACCENT_BLUE if col_idx == 1 else TEXT_MUTED)
        if col_idx == 0:
            p.font.bold = True

prs.save("Excel_Functions_Course_Pages_18_to_20.pptx")
print("Presentation generated successfully: Excel_Functions_Course_Pages_18_to_20.pptx")
